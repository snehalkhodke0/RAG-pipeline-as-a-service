"""Format-specific parsers that turn raw document bytes into a flat list of labeled
text sections. A "section" is the unit the chunker later splits independently: one PDF
page, or the text following one heading in a DOCX/Markdown/HTML document.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import List, Optional

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from markdown_it import MarkdownIt
from pypdf import PdfReader

from app.core.exceptions import ValidationAppError
from app.models.document import SourceType

_HEADING_STYLE_PREFIX = "Heading"
_HTML_HEADING_TAGS = ("h1", "h2", "h3", "h4", "h5", "h6")
_HTML_TEXT_TAGS = ("p", "li", "td", "blockquote", "pre")


@dataclass
class ParsedSection:
    text: str
    page_number: Optional[int]
    section_title: Optional[str]


@dataclass
class ParsedDocument:
    sections: List[ParsedSection]
    page_count: Optional[int]


def parse_pdf(content: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - pypdf raises various error types for corrupt files
        raise ValidationAppError(f"Could not parse PDF: {exc}") from exc

    sections: List[ParsedSection] = []
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(ParsedSection(text=text, page_number=index + 1, section_title=None))
    return ParsedDocument(sections=sections, page_count=len(reader.pages))


def parse_docx(content: bytes) -> ParsedDocument:
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - python-docx raises various error types
        raise ValidationAppError(f"Could not parse DOCX: {exc}") from exc

    sections: List[ParsedSection] = []
    current_title: Optional[str] = None
    current_lines: List[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            sections.append(ParsedSection(text=text, page_number=None, section_title=current_title))

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.startswith(_HEADING_STYLE_PREFIX):
            flush()
            current_title = text
            current_lines = []
        else:
            current_lines.append(text)
    flush()

    return ParsedDocument(sections=sections, page_count=None)


def parse_markdown(content: bytes) -> ParsedDocument:
    try:
        text_content = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValidationAppError(f"Markdown file is not valid UTF-8: {exc}") from exc

    tokens = MarkdownIt().parse(text_content)

    sections: List[ParsedSection] = []
    current_title: Optional[str] = None
    current_lines: List[str] = []
    awaiting_heading_text = False

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            sections.append(ParsedSection(text=text, page_number=None, section_title=current_title))

    for token in tokens:
        if token.type == "heading_open":
            flush()
            current_lines = []
            awaiting_heading_text = True
        elif token.type == "inline" and awaiting_heading_text:
            current_title = token.content.strip()
            awaiting_heading_text = False
        elif token.type == "inline":
            if token.content.strip():
                current_lines.append(token.content.strip())

    flush()
    return ParsedDocument(sections=sections, page_count=None)


def parse_html(content: bytes) -> ParsedDocument:
    try:
        soup = BeautifulSoup(content, "html.parser")
    except Exception as exc:  # noqa: BLE001 - BeautifulSoup can raise on malformed markup
        raise ValidationAppError(f"Could not parse HTML: {exc}") from exc

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    sections: List[ParsedSection] = []
    current_title: Optional[str] = None
    current_lines: List[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            sections.append(ParsedSection(text=text, page_number=None, section_title=current_title))

    body = soup.body or soup
    for element in body.find_all(_HTML_HEADING_TAGS + _HTML_TEXT_TAGS):
        text = element.get_text(separator=" ", strip=True)
        if not text:
            continue
        if element.name in _HTML_HEADING_TAGS:
            flush()
            current_title = text
            current_lines = []
        else:
            current_lines.append(text)
    flush()

    if not sections:
        # Fallback for pages with no recognized structural tags: grab all body text.
        fallback_text = body.get_text(separator="\n", strip=True)
        if fallback_text:
            sections.append(ParsedSection(text=fallback_text, page_number=None, section_title=None))

    return ParsedDocument(sections=sections, page_count=None)


_PARSERS = {
    SourceType.PDF: parse_pdf,
    SourceType.DOCX: parse_docx,
    SourceType.MARKDOWN: parse_markdown,
    SourceType.HTML: parse_html,
    SourceType.URL: parse_html,
}


def parse_document(source_type: SourceType, content: bytes) -> ParsedDocument:
    parser = _PARSERS.get(source_type)
    if parser is None:
        raise ValidationAppError(f"No parser registered for source type: {source_type}")
    return parser(content)
